"""Design endpoints — LaTeX template generation, resume assembly, and export."""

import asyncio
import io
import re

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from fastapi import APIRouter, HTTPException
from fastapi.responses import StreamingResponse
from pydantic import BaseModel

from core.ai_engine import (
    assemble_latex_resume,
    edit_latex_with_ai,
    generate_latex_template,
)

router = APIRouter()


# ── Schemas ───────────────────────────────────────────────────────────────────

class ProfileSummary(BaseModel):
    name: str = ""
    email: str = ""
    phone: str = ""
    location: str = ""
    linkedin: str = ""
    website: str = ""


class GenerateTemplateRequest(BaseModel):
    style_spec: str
    profile: ProfileSummary
    sections: list[str] = ["summary", "experience", "education", "skills"]


class GenerateTemplateResponse(BaseModel):
    latex: str


class EditTemplateRequest(BaseModel):
    latex: str
    instruction: str


class EditTemplateResponse(BaseModel):
    latex: str


class AssembleLatexRequest(BaseModel):
    latex_template: str
    resume_text: str
    profile: ProfileSummary


class AssembleLatexResponse(BaseModel):
    latex: str


# ── Endpoints ─────────────────────────────────────────────────────────────────

@router.post("/generate-template", response_model=GenerateTemplateResponse)
async def generate_template(req: GenerateTemplateRequest):
    """Generate a LaTeX resume template with lorem ipsum from a style description."""
    try:
        latex = await asyncio.to_thread(
            generate_latex_template,
            style_spec=req.style_spec,
            profile=req.profile.model_dump(),
            sections=req.sections,
        )
        return GenerateTemplateResponse(latex=latex)
    except Exception as exc:
        raise HTTPException(500, str(exc)) from exc


@router.post("/edit-template", response_model=EditTemplateResponse)
async def edit_template(req: EditTemplateRequest):
    """Apply a natural-language instruction to edit a LaTeX template."""
    try:
        latex = await asyncio.to_thread(edit_latex_with_ai, req.latex, req.instruction)
        return EditTemplateResponse(latex=latex)
    except Exception as exc:
        raise HTTPException(500, str(exc)) from exc


@router.post("/assemble-latex", response_model=AssembleLatexResponse)
async def assemble_latex(req: AssembleLatexRequest):
    """Replace lorem ipsum in a LaTeX template with real resume content."""
    try:
        latex = await asyncio.to_thread(
            assemble_latex_resume,
            latex_template=req.latex_template,
            resume_text=req.resume_text,
            profile=req.profile.model_dump(),
        )
        return AssembleLatexResponse(latex=latex)
    except Exception as exc:
        raise HTTPException(500, str(exc)) from exc


# ── DOCX Export ───────────────────────────────────────────────────────────────

class ExportDocxRequest(BaseModel):
    resume_text: str
    name: str = "resume"


def _build_docx(resume_text: str) -> bytes:
    """Convert plain-text resume to a formatted, ATS-friendly DOCX."""
    doc = Document()

    for section in doc.sections:
        section.top_margin    = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin   = Inches(1.0)
        section.right_margin  = Inches(1.0)

    # Remove the default empty paragraph
    for p in list(doc.paragraphs):
        p._element.getparent().remove(p._element)

    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    for raw in resume_text.split("\n"):
        stripped = raw.strip()

        if not stripped:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(2)
            continue

        # Bullet
        if re.match(r"^[•\-\*]\s", stripped):
            text = re.sub(r"^[•\-\*]\s*", "", stripped)
            p = doc.add_paragraph(style="List Bullet")
            run = p.add_run(text)
            run.font.size = Pt(10.5)
            p.paragraph_format.space_after = Pt(1)
            continue

        # Section header: ALL CAPS, no special chars
        if (stripped == stripped.upper() and len(stripped) > 2
                and not re.search(r"[|@\d]", stripped)):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after  = Pt(2)
            run = p.add_run(stripped)
            run.bold = True
            run.font.size = Pt(11)
            pPr = p._p.get_or_add_pPr()
            pBdr = OxmlElement("w:pBdr")
            bottom = OxmlElement("w:bottom")
            bottom.set(qn("w:val"), "single")
            bottom.set(qn("w:sz"), "4")
            bottom.set(qn("w:space"), "1")
            bottom.set(qn("w:color"), "AAAAAA")
            pBdr.append(bottom)
            pPr.append(pBdr)
            continue

        # Role | Company | Dates line
        if "|" in stripped:
            parts = [pt.strip() for pt in stripped.split("|")]
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(1)
            for j, part in enumerate(parts):
                run = p.add_run(part)
                run.font.size = Pt(10.5)
                run.bold = (j == 0)
                if j < len(parts) - 1:
                    sep = p.add_run("  |  ")
                    sep.font.size = Pt(10.5)
                    sep.font.color.rgb = RGBColor(0x99, 0x99, 0x99)
            continue

        # Default
        p = doc.add_paragraph()
        run = p.add_run(stripped)
        run.font.size = Pt(10.5)
        p.paragraph_format.space_after = Pt(1)

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.read()


@router.post("/export-docx")
async def export_docx(req: ExportDocxRequest):
    """Convert plain-text resume to a formatted DOCX and return for download."""
    try:
        docx_bytes = await asyncio.to_thread(_build_docx, req.resume_text)
        filename = re.sub(r"[^a-zA-Z0-9_-]", "_", req.name or "resume") + ".docx"
        return StreamingResponse(
            io.BytesIO(docx_bytes),
            media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            headers={"Content-Disposition": f'attachment; filename="{filename}"'},
        )
    except Exception as exc:
        raise HTTPException(500, str(exc)) from exc
